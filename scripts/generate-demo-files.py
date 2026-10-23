"""Create fictional multi-format demos locally (not committed to Git).

Usage: .venv/bin/python scripts/generate-demo-files.py
"""
from __future__ import annotations

import argparse
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

CONTENT = (
    "SHARESHIELD - FICTIONAL TEST DOCUMENT\n"
    "Email: demo@example.org\n"
    "Phone candidate: 9876543210\n"
    "PAN-format candidate: ABCDE1234F\n"
    "IFSC-format candidate: HDFC0001234\n"
    "Payment-card TEST number: 4111 1111 1111 1111\n"
    "UPI-format candidate: studentdemo@okaxis\n"
)


def sample_font():
    for path in ("/System/Library/Fonts/Supplemental/Arial.ttf",
                 "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if Path(path).is_file():
            return ImageFont.truetype(path, 28)
    return ImageFont.load_default()


def generate(out: Path):
    out.mkdir(parents=True, exist_ok=True)
    (out / "sample.txt").write_text(CONTENT, encoding="utf-8")
    (out / "sample.csv").write_text(
        "type,value\nemail,demo@example.org\nphone,9876543210\nPAN,ABCDE1234F\n",
        encoding="utf-8")

    img = Image.new("RGB", (1050, 300), "white")
    draw = ImageDraw.Draw(img)
    draw.text((35, 45), "Fictional example: demo@example.org", font=sample_font(), fill="#151515")
    draw.text((35, 120), "Phone candidate: 9876543210", font=sample_font(), fill="#151515")
    img.save(out / "sample.png", "PNG")
    img.save(out / "sample.jpg", "JPEG", quality=92)

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((65, 85), "SHARESHIELD SYNTHETIC DATA ONLY", fontsize=14)
    page.insert_text((65, 125), "Email: demo@example.org", fontsize=12)
    page.insert_text((65, 155), "Phone: 9876543210", fontsize=12)
    pdf.save(out / "sample-searchable.pdf")
    pdf.close()

    scanned = fitz.open()
    p = scanned.new_page(width=525, height=150)
    binary = BytesIO()
    img.save(binary, "PNG")
    p.insert_image(p.rect, stream=binary.getvalue())
    scanned.save(out / "sample-scanned.pdf")
    scanned.close()

    doc = Document()
    doc.add_heading("ShareShield Fictional Document", 1)
    for line in CONTENT.splitlines()[1:]:
        doc.add_paragraph(line)
    doc.save(out / "sample.docx")

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(4))
    box.text = CONTENT
    slides.save(out / "sample.pptx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fictional"
    sheet.append(["Type", "Value"])
    sheet.append(["Email", "demo@example.org"])
    sheet.append(["PAN", "ABCDE1234F"])
    sheet.append(["IFSC", "HDFC0001234"])
    workbook.save(out / "sample.xlsx")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", default="demo-output")
    args = parser.parse_args()
    folder = Path(args.output).expanduser().resolve()
    generate(folder)
    print(f"Generated fictional examples in {folder}; do not commit private files.")


if __name__ == "__main__":
    main()
