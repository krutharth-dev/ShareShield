"""Regression checks for spreadsheets, payment identifiers and audit integrity.

Only generated/synthetic values are used.
"""
import csv
import hashlib
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app import app
from backend.patterns import valid_luhn_card
from backend.review import evaluate
from backend.engine import Finding

client = TestClient(app)


def analyze(name: str, data: bytes):
    response = client.post("/analyze", files={"file": (name, data)})
    assert response.status_code == 200, response.text
    return response.json()


def redact(name: str, data: bytes, report: dict):
    response = client.post("/redact", files={"file": (name, data)},
                           data={"analysis_id": report["analysis_id"],
                                 "selected_ids": ",".join(str(f["id"]) for f in report["findings"])})
    assert response.status_code == 200, response.text
    return response


def test_csv_detection_redaction_and_audit_digest():
    original = "email,bank,card\nsample@example.org,HDFC0001234,4111 1111 1111 1111\n"
    data = original.encode()
    report = analyze("data.csv", data)
    assert {f["category"] for f in report["findings"]} >= {"email", "ifsc_candidate", "payment_card_candidate"}
    response = redact("data.csv", data, report)
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0] == ["email", "bank", "card"]
    assert rows[1] == ["[REDACTED]"] * 3
    assert response.headers["X-ShareShield-Verification"] == "PASS_SCOPED"
    audit = client.post("/verify",
                        files={"file": ("data.csv", data),
                               "redacted": ("shareshield-redacted.csv", response.content)},
                        data={"analysis_id": report["analysis_id"],
                              "selected_ids": ",".join(str(x["id"]) for x in report["findings"])})
    assert audit.status_code == 200, audit.text
    assert audit.json()["status"] == "PASS_SCOPED"
    assert audit.json()["output_sha256"] == hashlib.sha256(response.content).hexdigest()


def test_xlsx_redacts_text_and_number_cells_while_preserving_formulas():
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Email sample@example.org"
    sheet["B1"] = "IFSC HDFC0001234"
    sheet["C1"] = "=SUM(2,3)"
    sheet["D1"] = 4111111111111111
    extra = book.create_sheet("Internal")
    extra.sheet_state = "hidden"
    extra["A1"] = "PAN ABCDE1234F"
    output = io.BytesIO()
    book.save(output)
    data = output.getvalue()

    report = analyze("records.xlsx", data)
    assert len(report["findings"]) >= 4, report
    response = redact("records.xlsx", data, report)
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    saved = load_workbook(io.BytesIO(response.content))
    assert "[REDACTED]" in saved.active["A1"].value
    assert "[REDACTED]" in saved.active["B1"].value
    assert saved.active["C1"].value == "=SUM(2,3)"
    assert saved.active["D1"].value == "[REDACTED]"
    assert "[REDACTED]" in saved["Internal"]["A1"].value
    assert saved["Internal"].sheet_state == "hidden"


def test_payment_card_checksum_rejects_malformed_numbers():
    assert valid_luhn_card("4111 1111 1111 1111")
    assert not valid_luhn_card("4111111111111112")
    assert not valid_luhn_card("1111111111111111")
    response = analyze("test.txt", b"Invalid 4111111111111112")
    assert "payment_card_candidate" not in [f["category"] for f in response["findings"]]


def test_rejects_spoofed_xlsx_and_binary_csv():
    assert client.post("/analyze", files={"file": ("fake.xlsx", b"PKBAD")}).status_code in {415,422}
    assert client.post("/analyze", files={"file": ("binary.csv", b"abc\x00def")}).status_code == 422


def test_manual_also_requires_unverified_even_with_selected_detection():
    original = b"test@example.org"
    found = analyze("example.txt", original)
    result = redact("example.txt", original, found)
    selected = [Finding(category="email", value="test@example.org", page=1, rectangles=[], location="plain-text")]
    audit = evaluate(selected, result.content, ".txt", manual_count=1)
    assert audit["status"] == "UNVERIFIED"
    assert audit["selected_count"] == 1
