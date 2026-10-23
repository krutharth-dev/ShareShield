"""Synthetic-only integration tests for supported formats."""
import io
import shutil

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from backend.app import app

client = TestClient(app)


def request_analysis(name, content):
    response = client.post('/analyze', files={'file': (name, content)})
    assert response.status_code == 200, response.text
    return response.json()


def request_redaction(name, content, report, ids=None):
    if ids is None:
        ids=[f['id'] for f in report['findings']]
    return client.post('/redact', files={'file': (name, content)},
                       data={'analysis_id': report['analysis_id'],
                             'selected_ids': ','.join(map(str, ids))})


def pdf_sample():
    pdf=fitz.open(); page=pdf.new_page()
    page.insert_text((70,100),'Email: demo@example.org')
    page.insert_text((70,130),'Phone: 9876543210')
    result=pdf.tobytes();pdf.close();return result


def test_html_and_health():
    assert client.get('/').status_code==200
    assert 'Drop a file here' in client.get('/').text
    assert client.get('/health').json()['status']=='ok'


def test_searchable_pdf():
    data=pdf_sample(); report=request_analysis('test.pdf',data)
    assert [f['category'] for f in report['findings']]==['email','indian_mobile_candidate']
    r=request_redaction('test.pdf',data,report)
    assert r.status_code==200,r.text
    assert r.headers['X-ShareShield-Verification']=='PASS_SCOPED'
    with fitz.open(stream=r.content,filetype='pdf') as pdf:
        assert 'demo@example.org' not in pdf[0].get_text()
        assert '9876543210' not in pdf[0].get_text()


def test_docx():
    doc=Document(); doc.add_paragraph('Email demo@example.org')
    doc.add_table(rows=1,cols=1).cell(0,0).text='Phone 9876543210'
    buf=io.BytesIO(); doc.save(buf);data=buf.getvalue()
    report=request_analysis('test.docx',data)
    assert len(report['findings'])==2
    r=request_redaction('test.docx',data,report)
    assert r.status_code==200, r.text
    read=Document(io.BytesIO(r.content))
    assert '[REDACTED]' in read.paragraphs[0].text
    assert '[REDACTED]' in read.tables[0].cell(0,0).text


def test_pptx():
    p=Presentation();slide=p.slides.add_slide(p.slide_layouts[6]); box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1));box.text='demo@example.org'
    buf=io.BytesIO();p.save(buf);data=buf.getvalue()
    report=request_analysis('test.pptx',data)
    assert len(report['findings'])==1
    r=request_redaction('test.pptx',data,report)
    assert r.status_code==200,r.text
    assert '[REDACTED]' in Presentation(io.BytesIO(r.content)).slides[0].shapes[0].text


def test_txt_and_input_validation():
    data=b'demo@example.org\n'; report=request_analysis('input.txt',data)
    assert b'[REDACTED]' in request_redaction('input.txt',data,report).content
    assert client.post('/analyze',files={'file':('old.doc',b'a')}).status_code==415
    assert client.post('/analyze',files={'file':('photo.jpg',b'invalid')}).status_code==415
    assert request_redaction('input.txt',data,{'analysis_id':'wrong','findings':report['findings']}).status_code==409


@pytest.mark.skipif(not shutil.which('tesseract'), reason='system Tesseract missing')
def test_png_ocr_and_redaction():
    image=Image.new('RGB',(900,230),'white');draw=ImageDraw.Draw(image)
    try:font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',36)
    except OSError:font=ImageFont.load_default()
    draw.text((30,45),'Email: demo@example.org',fill='black',font=font)
    buf=io.BytesIO();image.save(buf,'PNG');data=buf.getvalue()
    report=request_analysis('image.png',data)
    assert any(f['value']=='demo@example.org' for f in report['findings']),report
    r=request_redaction('image.png',data,report)
    assert r.status_code==200,r.text
    assert r.headers['content-type'].startswith('image/png')
    assert r.content.startswith(b'\x89PNG')


@pytest.mark.skipif(not shutil.which('tesseract'), reason='system Tesseract missing')
def test_scanned_pdf_ocr():
    image=Image.new('RGB',(900,230),'white');draw=ImageDraw.Draw(image)
    try:font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',36)
    except OSError:font=ImageFont.load_default()
    draw.text((30,45),'demo@example.org',fill='black',font=font)
    raw=io.BytesIO();image.save(raw,'PNG')
    pdf=fitz.open();page=pdf.new_page(width=450,height=115);page.insert_image(fitz.Rect(0,0,450,115),stream=raw.getvalue());data=pdf.tobytes();pdf.close()
    report=request_analysis('scan.pdf',data)
    assert any(f['value']=='demo@example.org' for f in report['findings']),report
    response=request_redaction('scan.pdf',data,report)
    assert response.status_code==200,response.text
