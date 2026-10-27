import io
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from backend import engine, manual_text
from backend.app import app

client = TestClient(app)
SECRET = 'Redwood Project'


def fixture(ext):
    output = io.BytesIO()
    if ext == '.txt':
        return b'Public heading\nRedwood Project\nKeep this public sentence\n'
    if ext == '.csv':
        return b'public,private\nKeep this,Redwood Project\n'
    if ext == '.docx':
        doc = Document()
        doc.add_paragraph('Public heading')
        doc.add_paragraph(SECRET)
        doc.add_paragraph(SECRET)
        doc.save(output)
    elif ext == '.pptx':
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = SECRET
        doc.save(output)
    else:
        doc = Workbook()
        doc.active['A1'] = SECRET
        doc.active['A2'] = 'Public heading'
        doc.active['A3'] = '=1+1'
        doc.save(output)
    return output.getvalue()


@pytest.mark.parametrize('ext', ['.txt', '.csv', '.docx', '.pptx', '.xlsx'])
def test_manual_editor_and_exact_custom_redaction_round_trip(ext):
    original = fixture(ext)
    file = ('fictional' + ext, original)
    opened = client.post('/analyze', files={'file': file}, data={'manual_only': 'true'})
    assert opened.status_code == 200, opened.text
    analysis = opened.json()
    assert analysis['findings'] == []
    assert not analysis['preview_supported']
    preview = client.post('/text-preview', files={'file': file}, data={'analysis_id': analysis['analysis_id']})
    assert preview.status_code == 200, preview.text
    section = next(section for section in preview.json()['sections'] if SECRET in section['text'])
    selected = json.dumps([{'location': section['location'], 'value': SECRET}])
    params = {'analysis_id': analysis['analysis_id'], 'manual_text_regions': selected, 'redact_all': 'true'}
    output = client.post('/redact', files={'file': file}, data=params)
    assert output.status_code == 200, output.text
    text = '\n'.join(value for _, value in manual_text.sections(output.content, ext))
    assert SECRET not in text and '[REDACTED]' in text
    if ext == '.xlsx':
        workbook = load_workbook(io.BytesIO(output.content))
        assert workbook.active['A3'].value == '=1+1'
        workbook.close()
    verified = client.post('/verify', files={'file': file, 'redacted': ('redacted' + ext, output.content)}, data=params)
    assert verified.status_code == 200, verified.text
    assert verified.json()['status'] == 'UNVERIFIED'
    assert verified.json()['manual_text_count'] >= 1
    # Generic text does not match PII patterns, but a failed replacement must
    # still be detected by independent verification.
    unchanged = client.post('/verify', files={'file': file, 'redacted': file}, data=params)
    assert unchanged.status_code == 200
    assert unchanged.json()['status'] == 'FLAG'
    assert unchanged.json()['remaining_selected_count'] >= 1


def test_manual_only_bypasses_failed_ocr(monkeypatch):
    from PIL import Image
    source = io.BytesIO()
    Image.new('RGB', (50, 50), 'white').save(source, format='PNG')
    def broken(*args):
        raise AssertionError('Manual editor must not invoke OCR')
    monkeypatch.setattr(engine, 'inspect', broken)
    response = client.post('/analyze', files={'file': ('fictional.png', source.getvalue())}, data={'manual_only': 'true'})
    assert response.status_code == 200
    assert response.json()['preview_supported']


def test_selection_is_location_scoped_until_all_occurrences_requested():
    data = fixture('.docx')
    sections = manual_text.sections(data, '.docx')
    first = next(location for location, text in sections if text == SECRET)
    raw = json.dumps([{'location': first, 'value': SECRET}])
    chosen = manual_text.parse(raw, data, '.docx', False)
    output, _, _ = engine.redact(data, '.docx', chosen)
    assert sum(SECRET in text for _, text in manual_text.sections(output, '.docx')) == 1
    assert len(manual_text.parse(raw, data, '.docx', True)) == 2


@pytest.mark.parametrize('raw', ['null', '{}', '[{}]', '[{"location":"missing","value":"Redwood"}]', '[{"location":"plain-text","value":""}]', '[{"location":"plain-text","value":"absent"}]'])
def test_invalid_or_fabricated_text_selections_rejected(raw):
    opened = client.post('/analyze', files={'file': ('fictional.txt', fixture('.txt'))}).json()
    response = client.post('/redact', files={'file': ('fictional.txt', fixture('.txt'))},
        data={'analysis_id': opened['analysis_id'], 'manual_text_regions': raw})
    assert response.status_code == 422


def test_text_preview_is_bound_to_reviewed_source_and_limited():
    opened = client.post('/analyze', files={'file': ('fictional.txt', fixture('.txt'))}).json()
    changed = client.post('/text-preview', files={'file': ('fictional.txt', b'changed')}, data={'analysis_id': opened['analysis_id']})
    assert changed.status_code == 409
    large = b'a' * 200_001
    opened = client.post('/analyze', files={'file': ('fictional.txt', large)}, data={'manual_only': 'true'}).json()
    response = client.post('/text-preview', files={'file': ('fictional.txt', large)}, data={'analysis_id': opened['analysis_id']})
    assert response.status_code == 413
