"""Synthetic PDF and Office release smoke tests.

Runs in the Docker image build, using only invented sample data.
"""
import io

import fitz
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from backend import engine, review


TEST_EMAIL = "fictional.demo@example.org"


def main() -> None:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 80), f"Fictional email: {TEST_EMAIL}", fontsize=12)
    pdf_bytes = pdf.tobytes()
    pdf.close()

    findings, _ = engine.inspect(pdf_bytes, ".pdf")
    chosen = [f for f in findings if f.value == TEST_EMAIL]
    assert chosen and chosen[0].rectangles, "Native PDF detection and bounding boxes"
    output, _, _ = review.redact_visual(pdf_bytes, ".pdf", chosen, [], "blackout")
    with engine.open_pdf(output) as inspected:
        assert TEST_EMAIL not in inspected[0].get_text(), "PDF sensitive text not removed"
    remaining, _ = engine.inspect(output, ".pdf")
    assert not any(f.value == TEST_EMAIL for f in remaining), "PDF still detects sensitive text"
    report = review.evaluate(chosen, output, ".pdf")
    assert report["status"] in {"PASS_SCOPED", "UNVERIFIED"}, report["status"]

    word = Document()
    word.add_paragraph(f"Email: {TEST_EMAIL}")
    buffer = io.BytesIO()
    word.save(buffer)
    matches, _ = engine.inspect(buffer.getvalue(), ".docx")
    assert any(f.value == TEST_EMAIL for f in matches), "DOCX parsing failed"

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    shape.text = f"Email: {TEST_EMAIL}"
    buffer = io.BytesIO()
    slides.save(buffer)
    matches, _ = engine.inspect(buffer.getvalue(), ".pptx")
    assert any(f.value == TEST_EMAIL for f in matches), "PPTX parsing failed"

    # Call the actual ASGI endpoint functions with in-memory UploadFiles.
    # This verifies async upload decoding and threaded document processing.
    import asyncio
    from fastapi import UploadFile
    from backend.app import analyze, redact_selected, verify_result

    async def endpoint_roundtrip() -> None:
        def upload(data: bytes, filename: str) -> UploadFile:
            return UploadFile(file=io.BytesIO(data), filename=filename)

        analysis = await analyze(upload(pdf_bytes, "fictional.pdf"))
        ids = [f["id"] for f in analysis["findings"] if f["value"] == TEST_EMAIL]
        assert ids, "Analyze endpoint did not return PDF finding"
        redacted_response = await redact_selected(
            upload(pdf_bytes, "fictional.pdf"),
            analysis_id=analysis["analysis_id"],
            selected_ids=",".join(str(i) for i in ids),
            manual_regions="[]",
            redact_all=False,
            redaction_mode="blackout",
        )
        assert redacted_response.body.startswith(b"%PDF-"), "Redact endpoint did not output PDF"
        verification = await verify_result(
            upload(pdf_bytes, "fictional.pdf"),
            upload(redacted_response.body, "fictional-redacted.pdf"),
            analysis_id=analysis["analysis_id"],
            selected_ids=",".join(str(i) for i in ids),
            manual_count=0,
            redact_all=False,
            redaction_mode="blackout",
        )
        assert verification["status"] in {"PASS_SCOPED", "UNVERIFIED"}, verification["status"]

    asyncio.run(endpoint_roundtrip())
    print("RELEASE CHECK PASSED: PDF analyze/redact/verify API, DOCX scan, PPTX scan")


if __name__ == "__main__":
    main()
